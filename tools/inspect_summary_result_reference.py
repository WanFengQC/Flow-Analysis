"""输出 SUMMARY_result 参考工作簿的结构与样式诊断信息。"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _color_value(color: Any) -> str | None:
    """将 openpyxl 颜色对象转换为可比较的 RGB 或索引值。"""

    if color is None:
        return None
    return getattr(color, "rgb", None) or getattr(color, "indexed", None)


def inspect_summary_result(workbook_path: Path) -> dict[str, Any]:
    """读取用户参考工作簿，不修改原文件并返回关键格式信息。"""

    workbook = load_workbook(workbook_path, data_only=False)
    worksheet = workbook["SUMMARY_result"]
    sample_cells = (
        "A1",
        "B1",
        "A2",
        "B2",
        "A3",
        "B3",
        "A16",
        "B16",
        "A17",
        "B17",
        "A18",
        "D18",
        "E18",
        "G18",
        "H18",
        "J18",
    )
    styles: dict[str, dict[str, Any]] = {}
    for coordinate in sample_cells:
        cell = worksheet[coordinate]
        styles[coordinate] = {
            "styleId": cell.style_id,
            "fillRgb": _color_value(cell.fill.fgColor),
            "numberFormat": cell.number_format,
            "font": {
                "name": cell.font.name,
                "size": cell.font.sz,
                "bold": cell.font.bold,
            },
            "alignment": {
                "horizontal": cell.alignment.horizontal,
                "vertical": cell.alignment.vertical,
                "wrapText": cell.alignment.wrap_text,
            },
            "border": {
                "left": cell.border.left.style,
                "right": cell.border.right.style,
                "top": cell.border.top.style,
                "bottom": cell.border.bottom.style,
            },
        }

    return {
        "sheetNames": workbook.sheetnames,
        "summaryResult": {
            "maxRow": worksheet.max_row,
            "maxColumn": worksheet.max_column,
            "defaultRowHeight": worksheet.sheet_format.defaultRowHeight,
            "columnWidths": {
                column: dimension.width
                for column, dimension in worksheet.column_dimensions.items()
                if dimension.width is not None
            },
            "mergedCells": [str(item) for item in worksheet.merged_cells.ranges],
            "freezePanes": worksheet.freeze_panes,
            "autoFilter": worksheet.auto_filter.ref,
            "showGridLines": worksheet.sheet_view.showGridLines,
            "sampleStyles": styles,
        },
    }


def main() -> int:
    """提供可复现的命令行检查入口，便于更新参考模板时复核。"""

    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    arguments = parser.parse_args()
    print(
        json.dumps(
            inspect_summary_result(arguments.workbook),
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
