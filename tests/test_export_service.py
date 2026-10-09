"""SUMMARY_result Excel 导出的结构、样式和数据隔离测试。"""

from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from copy import copy
from unittest.mock import patch

from openpyxl import load_workbook

from services.export_service import ExportService, ExportServiceError
from services.excel_export.template_support import display_ratio


class ExportServiceTest(unittest.TestCase):
    """验证导出复刻模板布局，且只使用传入 Final Analysis 数据。"""

    @classmethod
    def setUpClass(cls) -> None:
        """完整模板较大，所有只读断言复用同一份已导出工作簿。"""

        cls.service = ExportService()
        cls.rows = cls._build_rows()
        cls.output_directory = TemporaryDirectory()
        cls.output = Path(cls.output_directory.name) / "analysis.xlsx"
        cls.service.export_analysis(rows=cls.rows, output_path=cls.output)
        cls.template_workbook = load_workbook(
            cls.service._template_path(),
            data_only=False,
        )
        cls.output_workbook = load_workbook(cls.output, data_only=False)

    @classmethod
    def tearDownClass(cls) -> None:
        """清理由测试生成的临时工作簿。"""

        cls.template_workbook.close()
        cls.output_workbook.close()
        cls.output_directory.cleanup()

    @staticmethod
    def _build_rows() -> list[dict[str, object]]:
        """构造覆盖两个自然月、已打标和未完成状态的固定 FinalAnalysisRows。"""

        return [
            {
                "month": "202608", "word": "weighted", "frequency": 2,
                "weight": Decimal("12000"), "total": 200.0, "ratio": 0.06,
                "naturalRatio": 0.4, "adRatio": 0.6,
                "topPhrases": ["weighted stuffed animal", "weighted plush"],
                "label": "1核心词", "labelReason": "核心商品词",
                "labelSource": "AI共识", "sourceAsinStats": {"B0TEST": {"exposure": 12.5}},
            },
            {
                "month": "202608", "word": "soft", "frequency": 1,
                "weight": Decimal("3000"), "total": 80.0, "ratio": 0.04,
                "naturalRatio": None, "adRatio": 0.2,
                "topPhrases": ["soft stuffed animal"],
                "label": "3属性", "labelReason": "触感属性",
            },
            {
                "month": "202607", "word": "weighted", "frequency": 3,
                "weight": Decimal("8000"), "total": 160.0, "ratio": 0.05,
                "naturalRatio": 0.5, "adRatio": 0.5,
                "topPhrases": ["weighted animal"],
                "label": "1核心词", "labelReason": "核心商品词",
            },
            {
                "month": "202607", "word": "old core", "frequency": 1,
                "weight": Decimal("1000"), "total": 20.0, "ratio": 0.02,
                "naturalRatio": 0.1, "adRatio": 0.9,
                "topPhrases": [], "label": "1核心词", "labelReason": "次级核心词",
            },
            {
                "month": "202607", "word": "unreviewed", "frequency": 1,
                "weight": Decimal("500"), "total": 5.0, "ratio": None,
                "naturalRatio": None, "adRatio": None,
                "topPhrases": [], "label": None, "labelSource": "INCOMPLETE",
                "labelReason": "不得导出",
            },
        ]

    def test_export_writes_summary_result_with_month_blocks_and_reference_styles(self):
        """月份横向并排、字段固定、颜色和边框必须来自参考模板。"""

        workbook = self.output_workbook
        template_workbook = self.template_workbook
        self.assertEqual(workbook.sheetnames, template_workbook.sheetnames)
        worksheet = workbook["SUMMARY_result"]

        self.assertEqual(worksheet["A1"].value, "8月")
        self.assertEqual(worksheet["L1"].value, "7月")
        self.assertEqual(worksheet["K1"].value, None)
        self.assertEqual(worksheet["A2"].value, "打标")
        self.assertEqual(worksheet["G2"].value, "打标weight占比")
        self.assertEqual(worksheet["A16"].value, "1核心词")
        self.assertEqual(worksheet["L16"].value, "1核心词")
        self.assertEqual(worksheet["A17"].value, "单词")
        self.assertEqual(worksheet["J17"].value, "导入短语数据")
        self.assertEqual(worksheet["A18"].value, "weighted")
        self.assertEqual(worksheet["J18"].value, "weighted stuffed animal | weighted plush")
        self.assertEqual(worksheet["G18"].number_format, "0.00%")
        self.assertEqual(worksheet["G18"].fill.fgColor.rgb[-6:], "FFEBAD")
        self.assertEqual(worksheet["E1"].number_format, "0.00%")
        self.assertEqual(worksheet["B1"].number_format, '0.00"万"')
        self.assertEqual(worksheet["A1"].fill.fgColor.rgb[-6:], "CFE2F3")
        self.assertEqual(worksheet["A2"].fill.fgColor.rgb[-6:], "F4CCCC")
        self.assertEqual(worksheet["A1"].border.left.style, "medium")
        self.assertEqual(worksheet["G1"].border.right.style, "medium")
        self.assertEqual(worksheet.freeze_panes, None)
        self.assertEqual(worksheet.auto_filter.ref, None)
        self.assertEqual(worksheet.sheet_view.showGridLines, None)
        self.assertTrue(
            all(
                worksheet.column_dimensions[column].width == 10.0
                for column in "ABCDEFGHIJKLMNOPQRSTU"
            )
        )
        self.assertEqual(worksheet.sheet_format.defaultRowHeight, 15.0)
        template = template_workbook["SUMMARY_result"]
        for coordinate in (
            "A1", "B1", "D1", "A2", "B2", "A3", "B3", "D3", "G3",
            "A16", "B16", "E16", "G16", "J16", "A17", "B17", "J17",
            "A18", "B18", "D18", "E18", "H18", "J18",
        ):
            # 导出会按月份区块补齐外框；该边框是运行时结构的一部分，已在
            # 上方分别断言。其余样式必须与模板一致。
            self._assert_template_style(
                template[coordinate],
                worksheet[coordinate],
                compare_border=False,
            )

    def test_export_accepts_recent_30_days_sentinel_as_an_independent_range(self):
        """空字符串仅在 API 内部表示最近30天，导出必须保留并明确展示。"""

        recent_rows = [
            {
                "month": "",
                "word": "weighted",
                "frequency": 3,
                "weight": 100.0,
                "total": 200.0,
                "ratio": 0.5,
                "naturalRatio": 0.4,
                "adRatio": 0.6,
                "topPhrases": ["weighted stuffed animals"],
                "label": "1核心词",
            }
        ]
        with TemporaryDirectory() as directory:
            output_path = Path(directory) / "recent-30-days.xlsx"
            ExportService().export_analysis(
                rows=recent_rows,
                output_path=output_path,
            )
            workbook = load_workbook(output_path, data_only=False)
            try:
                worksheet = workbook["SUMMARY_result"]
                self.assertEqual(worksheet["A1"].value, "最近30天")
                self.assertEqual(worksheet["A18"].value, "weighted")
            finally:
                workbook.close()

    def test_unlabeled_rows_remain_visible_in_invalid_section_with_blank_tag_and_mark(self):
        """DISAGREEMENT 或 INCOMPLETE 不能伪造标签，但词必须保留。"""

        worksheet = self.output_workbook["SUMMARY_result"]
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value != "unreviewed":
                    continue
                self.assertIsNone(worksheet.cell(cell.row, cell.column + 2).value)
                self.assertIsNone(worksheet.cell(cell.row, cell.column + 3).value)
                break
            else:
                continue
            break
        else:
            self.fail("未在无效词区块找到未完成词")
        values = [cell.value for row in worksheet.iter_rows() for cell in row]
        self.assertNotIn("标签来源", values)
        self.assertNotIn("匹配搜索词数", values)
        self.assertNotIn("Product Context", values)

    def test_export_preserves_all_template_sheets_and_clears_unimplemented_content(self):
        """完整工作簿必须保留全部 Sheet 结构，未接入 Sheet 不得泄漏旧数据。"""

        output_workbook = self.output_workbook
        template_workbook = self.template_workbook
        self.assertEqual(output_workbook.sheetnames, template_workbook.sheetnames)
        self.assertEqual(len(output_workbook.worksheets), len(template_workbook.worksheets))
        for template_sheet, output_sheet in zip(
            template_workbook.worksheets,
            output_workbook.worksheets,
        ):
            self.assertEqual(output_sheet.title, template_sheet.title)
            self.assertEqual(output_sheet.sheet_state, template_sheet.sheet_state)
            self.assertEqual(
                tuple(output_sheet.merged_cells.ranges),
                tuple(template_sheet.merged_cells.ranges),
            )
            self.assertEqual(output_sheet.freeze_panes, template_sheet.freeze_panes)
            self.assertEqual(output_sheet.auto_filter.ref, template_sheet.auto_filter.ref)
            self.assertEqual(output_sheet.page_margins, template_sheet.page_margins)
            self.assertEqual(output_sheet.page_setup, template_sheet.page_setup)
            if output_sheet.title == "SUMMARY_result":
                continue
            self.assertEqual(output_sheet.max_row, template_sheet.max_row)
            self.assertEqual(output_sheet.max_column, template_sheet.max_column)
            self.assertEqual(
                output_sheet.column_dimensions["A"].width,
                template_sheet.column_dimensions["A"].width,
            )
            self._assert_template_style(template_sheet["A1"], output_sheet["A1"])
            self.assertFalse(
                any(cell.value is not None for cell in output_sheet._cells.values())
            )

    def test_export_populates_all_runtime_sheets_and_replaces_template_asin(self):
        """单 ASIN 导出必须保留五个真实 Sheet，且不得泄漏模板示例 ASIN。"""

        with TemporaryDirectory() as directory:
            output = Path(directory) / "full_runtime.xlsx"
            self.service.export_analysis(
                rows=self.rows,
                output_path=output,
                analysis_raw=self._build_analysis_raw(),
                reversing_results={
                    "202608": {"B0NEWASIN": {"items": []}},
                    "202607": {"B0NEWASIN": {"items": []}},
                },
                analysis_word_results=self._build_word_results(),
                analysis_asin_word_results=self._build_asin_word_results(),
            )
            workbook = load_workbook(output, data_only=False)
            self.addCleanup(workbook.close)
            self.assertEqual(
                workbook.sheetnames,
                [
                    "SUMMARY_raw", "SUMMARY_result", "SUMMARY_result_list",
                    "B0NEWASIN_result", "B0NEWASIN_raw",
                ],
            )
            self.assertEqual(workbook["SUMMARY_raw"]["A1"].value, "8月")
            self.assertEqual(workbook["SUMMARY_raw"]["A3"].value, "alpha")
            self.assertEqual(workbook["SUMMARY_raw"]["B3"].value, 1200)
            self.assertEqual(workbook["SUMMARY_raw"]["I3"].value, "B0NEWASIN")
            self.assertEqual(workbook["SUMMARY_raw"]["J2"].value, "guestId")
            self.assertEqual(workbook["SUMMARY_raw"]["J3"].value, "guest-1")
            self.assertEqual(workbook["SUMMARY_raw"]["K2"].value, "station")
            self._assert_template_style(
                self.template_workbook["SUMMARY_raw"]["K2"],
                workbook["SUMMARY_raw"]["J2"],
            )
            self._assert_template_style(
                self.template_workbook["SUMMARY_raw"]["K3"],
                workbook["SUMMARY_raw"]["J3"],
            )
            self.assertNotIn(
                "source_parent_asin",
                {
                    cell.value
                    for row in workbook["SUMMARY_raw"].iter_rows(
                        min_row=2,
                        max_row=2,
                    )
                    for cell in row
                },
            )
            self.assertEqual(
                workbook["SUMMARY_raw"].cell(3, 48).value,
                '["NATURAL_SEARCHING", "SPONSOR_BRAND", "ADS"]',
            )
            result_list = workbook["SUMMARY_result_list"]
            self.assertEqual(result_list["A2"].value, "时间")
            self.assertEqual(result_list["B2"].value, "父ASIN")
            self.assertEqual(result_list["C2"].value, "单词")
            self.assertEqual(result_list["A3"].value, "8月")
            self.assertEqual(result_list["B3"].value, "B0NEWASIN")
            self.assertEqual(result_list["C3"].value, "weighted")
            self.assertEqual(result_list["D3"].value, 2)
            self.assertEqual(result_list["E3"].value, "1核心词")
            self.assertEqual(result_list["G3"].value, 12000)
            self._assert_template_style(
                self.template_workbook["SUMMARY_result_list"]["A1"],
                result_list["A1"],
            )
            self._assert_template_style(
                self.template_workbook["SUMMARY_result_list"]["A2"],
                result_list["A2"],
            )
            self._assert_template_style(
                self.template_workbook["SUMMARY_result_list"]["A3"],
                result_list["A3"],
            )
            self.assertTrue(
                {"关键词", "来源 ASIN", "预估周曝光量", "点击量"}.isdisjoint(
                    {cell.value for cell in result_list[2]}
                )
            )
            self.assertTrue(
                {"单词", "频次", "Weight", "Total"}.issubset(
                    {cell.value for cell in result_list[2]}
                )
            )
            self.assertEqual(
                sum(
                    1
                    for row in result_list.iter_rows(min_row=3)
                    for cell in row
                    if cell.value in {"weighted", "soft", "old core", "unreviewed"}
                ),
                sum(len(rows) for rows in self._build_word_results().values()),
            )
            expected_words = {
                row["word"]
                for rows in self._build_word_results().values()
                for row in rows
            }
            list_words = {
                cell.value
                for row in result_list.iter_rows(min_row=3)
                for cell in row
                if cell.value in expected_words
            }
            summary_words = {
                cell.value
                for row in workbook["SUMMARY_result"].iter_rows()
                for cell in row
                if cell.value in expected_words
            }
            self.assertEqual(list_words, summary_words)
            asin_result = workbook["B0NEWASIN_result"]
            self.assertEqual(asin_result["A2"].value, "单词")
            self.assertEqual(asin_result["A3"].value, "weighted")
            self.assertEqual(asin_result["B3"].value, 2)
            self.assertEqual(asin_result["E3"].value, 12000)
            self._assert_template_style(
                self.template_workbook["B0TMPL0001_result"]["A2"],
                asin_result["A2"],
            )
            self._assert_template_style(
                self.template_workbook["B0TMPL0001_result"]["A3"],
                asin_result["A3"],
            )
            self.assertEqual(result_list["C3"].value, asin_result["A3"].value)
            self.assertEqual(result_list["G3"].value, asin_result["E3"].value)
            self.assertTrue(
                {"关键词", "来源 ASIN", "预估周曝光量", "点击量"}.isdisjoint(
                    {cell.value for cell in asin_result[2]}
                )
            )
            self.assertEqual(
                sum(
                    1
                    for row in asin_result.iter_rows(min_row=3)
                    for cell in row
                    if cell.value in {"weighted", "soft", "old core", "unreviewed"}
                ),
                sum(
                    len(asins["B0NEWASIN"])
                    for asins in self._build_asin_word_results().values()
                ),
            )
            self.assertEqual(workbook["B0NEWASIN_raw"]["A3"].value, "alpha")
            self.assertEqual(
                workbook["SUMMARY_raw"]["A1"].fill.fgColor.rgb,
                self.template_workbook["SUMMARY_raw"]["A1"].fill.fgColor.rgb,
            )
            self.assertEqual(
                workbook["B0NEWASIN_result"]["A1"].fill.fgColor.rgb,
                self.template_workbook["B0TMPL0001_result"]["A1"].fill.fgColor.rgb,
            )
            self.assertNotIn("B0TMPL0001", workbook.sheetnames)
            self.assertFalse(
                any(
                    cell.value == "B0TMPL0001"
                    for worksheet in workbook.worksheets
                    for cell in worksheet._cells.values()
                )
            )

    def test_export_creates_a_sheet_pair_for_each_current_asin(self):
        """多 ASIN 时新增 Sheet 必须由同一示例模板复制，且名称不固定。"""

        with TemporaryDirectory() as directory:
            output = Path(directory) / "multi_asin.xlsx"
            raw = self._build_analysis_raw()
            raw["202608"].append({
                **raw["202608"][0], "source_asin": "B0SECOND", "keywords": "beta",
            })
            asin_word_results = self._build_asin_word_results()
            asin_word_results["202608"]["B0SECOND"] = [{
                "word": "beta", "frequency": 1, "weight": 900,
                "total": 100, "ratio": 0.09, "naturalRatio": 0.5,
                "adRatio": 0.5, "topPhrases": ["beta stuffed animal"],
            }]
            self.service.export_analysis(
                rows=self.rows,
                output_path=output,
                analysis_raw=raw,
                reversing_results={"202608": {"B0NEWASIN": {}, "B0SECOND": {}}},
                analysis_word_results=self._build_word_results(),
                analysis_asin_word_results=asin_word_results,
            )
            workbook = load_workbook(output, data_only=False)
            self.addCleanup(workbook.close)
            self.assertIn("B0NEWASIN_result", workbook.sheetnames)
            self.assertIn("B0NEWASIN_raw", workbook.sheetnames)
            self.assertIn("B0SECOND_result", workbook.sheetnames)
            self.assertIn("B0SECOND_raw", workbook.sheetnames)
            self.assertEqual(workbook["B0SECOND_result"]["A3"].value, "beta")
            self.assertEqual(workbook["B0SECOND_raw"]["A3"].value, "beta")

    def test_export_reports_missing_asin_word_results_without_keyword_result_fallback(self):
        """ASIN 词结果缺失时不能把关键词 RESULT 冒充为 Word Analysis。"""

        with TemporaryDirectory() as directory:
            output = Path(directory) / "missing_asin_words.xlsx"
            with self.assertRaisesRegex(
                ExportServiceError,
                "缺少 ASIN Word Analysis 数据：B0NEWASIN",
            ):
                self.service.export_analysis(
                    rows=self.rows,
                    output_path=output,
                    analysis_raw=self._build_analysis_raw(),
                    reversing_results={"202608": {"B0NEWASIN": {}}},
                    analysis_word_results=self._build_word_results(),
                    analysis_asin_word_results={},
                )

    @staticmethod
    def _build_analysis_raw() -> dict[str, list[dict[str, object]]]:
        """构造已完成分析阶段的 RAW 快照，模拟导出线程只读输入。"""

        return {
            "202608": [{
                "month": "202608", "source_asin": "B0NEWASIN", "keywords": "alpha",
                "calculatedWeeklySearches": 1200, "searchesRank": 88, "searches": 3400,
                "impressions": 5000, "clicks": 200, "naturalRatio": 0.7, "adRatio": 0.3,
                "station": "US", "keywordCn": "阿尔法", "guestId": "guest-1",
                "trafficKeywordTypes": ["NATURAL_SEARCHING", "SPONSOR_BRAND", "ADS"],
            }],
            "202607": [{
                "month": "202607", "source_asin": "B0NEWASIN", "keywords": "gamma",
                "calculatedWeeklySearches": 900, "searchesRank": 91, "searches": 2800,
                "impressions": 4000, "clicks": 160, "naturalRatio": 0.6, "adRatio": 0.4,
                "station": "US", "keywordCn": "伽马",
            }],
        }

    @staticmethod
    def _build_word_results() -> dict[str, list[dict[str, object]]]:
        """构造正式 Word Result，字段与 SUMMARY_result_list 的原表头一致。"""

        return {
            "202608": [
                {
                    "word": "weighted", "frequency": 2, "weight": 12000,
                    "total": 200, "ratio": 0.06, "naturalRatio": 0.4,
                    "adRatio": 0.6, "topPhrases": ["weighted stuffed animal"],
                    "sourceAsinStats": {"B0NEWASIN": {"exposure": 1200}},
                },
                {
                    "word": "soft", "frequency": 1, "weight": 3000,
                    "total": 80, "ratio": 0.04, "naturalRatio": None,
                    "adRatio": 0.2, "topPhrases": ["soft stuffed animal"],
                    "sourceAsinStats": {"B0NEWASIN": {"exposure": 300}},
                },
            ],
            "202607": [
                {
                    "word": "weighted", "frequency": 3, "weight": 8000,
                    "total": 160, "ratio": 0.05, "naturalRatio": 0.5,
                    "adRatio": 0.5, "topPhrases": ["weighted animal"],
                    "sourceAsinStats": {"B0NEWASIN": {"exposure": 800}},
                },
                {
                    "word": "old core", "frequency": 1, "weight": 1000,
                    "total": 20, "ratio": 0.02, "naturalRatio": 0.1,
                    "adRatio": 0.9, "topPhrases": [],
                    "sourceAsinStats": {"B0NEWASIN": {"exposure": 100}},
                },
                {
                    "word": "unreviewed", "frequency": 1, "weight": 500,
                    "total": 5, "ratio": None, "naturalRatio": None,
                    "adRatio": None, "topPhrases": [],
                    "sourceAsinStats": {"B0NEWASIN": {"exposure": 50}},
                },
            ],
        }

    @classmethod
    def _build_asin_word_results(cls) -> dict[str, dict[str, list[dict[str, object]]]]:
        """构造分析阶段生成的 ASIN 独立 Word Result 快照。"""

        return {
            month: {"B0NEWASIN": [dict(row) for row in rows]}
            for month, rows in cls._build_word_results().items()
        }

    def test_template_asset_is_sanitized_and_contains_no_reference_business_data(self):
        """模板只能携带结构和样式，不得把参考工作簿业务数据带入发布物。"""

        template = self.service._template_path()
        self.assertTrue(template.is_file())
        workbook = self.template_workbook
        self.assertEqual(
            workbook.sheetnames,
            [
                "SUMMARY_raw",
                "SUMMARY_result",
                "SUMMARY_result_list",
                "B0TMPL0001_result",
                "B0TMPL0001_raw",
            ],
        )
        self.assertFalse(
            any(
                cell.value is not None
                for worksheet in workbook.worksheets
                for cell in worksheet._cells.values()
            )
        )

    def test_empty_final_rows_are_rejected_without_writing_file(self):
        """没有 Final Dataset 时不能伪造空导出文件。"""

        with TemporaryDirectory() as directory:
            output = Path(directory) / "empty.xlsx"
            with self.assertRaises(ExportServiceError):
                self.service.export_analysis(rows=[], output_path=output)
            self.assertFalse(output.exists())

    def test_zero_total_ratio_is_written_as_zero(self):
        """SUMMARY_result 及两个词明细 Sheet 都不应把确定零值导出为空白。"""

        self.assertEqual(ExportService._safe_divide(0, 0), 0.0)
        self.assertEqual(display_ratio(None, 0), 0.0)
        self.assertIsNone(display_ratio(None, None))

    def test_save_failure_removes_partial_output_file(self):
        """保存异常不能留下正式 xlsx 或临时 part 文件。"""

        with TemporaryDirectory() as directory:
            output = Path(directory) / "failed.xlsx"
            workbook = load_workbook(self.service._template_path())
            self.addCleanup(workbook.close)
            with (
                patch(
                    "services.export_service.load_workbook",
                    return_value=workbook,
                ),
                patch.object(
                    workbook,
                    "save",
                    side_effect=OSError("disk unavailable"),
                ),
            ):
                with self.assertRaisesRegex(ExportServiceError, "Excel 导出失败"):
                    self.service.export_analysis(rows=self.rows, output_path=output)
            self.assertFalse(output.exists())
            self.assertFalse((Path(directory) / ".failed.xlsx.part").exists())

    def _assert_template_style(
        self,
        expected,
        actual,
        *,
        compare_border: bool = True,
    ) -> None:
        """对照清洗后的参考模板，验证字体、填充、边框、对齐和数字格式。"""

        self.assertEqual(copy(actual.font), copy(expected.font))
        self.assertEqual(copy(actual.fill), copy(expected.fill))
        if compare_border:
            self.assertEqual(copy(actual.border), copy(expected.border))
        self.assertEqual(copy(actual.alignment), copy(expected.alignment))
        self.assertEqual(actual.number_format, expected.number_format)


if __name__ == "__main__":
    unittest.main()
